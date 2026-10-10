"""Inspect a local MODMA EEG dataset without modifying its files."""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "src"))

from snn_depression.common.paths import DATA_ROOT


EXPECTED_MD5 = "7f5ea8c89c550443dc740c5c9e9d3867"
TABLE_PATTERN = re.compile(r"subjects?.*information|demograph|phq", re.IGNORECASE)


def find_dataset_root(data_root: Path, requested: Path | None) -> Path:
    """Find the requested dataset directory or the MODMA dataset in DATA_ROOT."""
    if requested is not None:
        if not requested.is_dir():
            raise FileNotFoundError(f"Dataset directory does not exist: {requested}")
        return requested

    expected = data_root / "raw" / "modma"
    if expected.is_dir():
        return expected

    named = data_root / "EEG_128channels_resting_lanzhou_2015"
    if named.is_dir():
        return named

    candidates = []
    for path in data_root.rglob("*"):
        if path.is_dir() and any(
            TABLE_PATTERN.search(file.name)
            and file.suffix.lower() in {".xlsx", ".xls", ".csv"}
            for file in path.iterdir()
            if file.is_file()
        ):
            candidates.append(path)
    if candidates:
        return candidates[0]
    raise FileNotFoundError(
        f"No MODMA folder or demographics table found under DATA_ROOT={data_root}"
    )


def md5_file(path: Path) -> str:
    """Return a file's MD5 digest in bounded-memory chunks."""
    digest = hashlib.md5()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def print_file_inventory(root: Path, files: list[Path]) -> None:
    """Print extensions, relative paths, and filename ID patterns."""
    formats = Counter(path.suffix.lower() or "<no extension>" for path in files)
    print("\nFILE FORMATS")
    for extension, count in sorted(formats.items()):
        print(f"  {extension}: {count}")

    print("\nFILE NAMING (up to 60 files)")
    ids: list[str] = []
    for path in files[:60]:
        relative = path.relative_to(root)
        match = re.match(r"(\d+)", path.stem)
        if match:
            ids.append(match.group(1))
        print(f"  {relative}")
    prefixes = Counter(subject_id[:4] for subject_id in ids)
    print(f"  Leading numeric filename IDs: {len(ids)} shown; 4-digit prefix counts={dict(prefixes)}")


def print_archive_checks(root: Path, files: list[Path]) -> None:
    """Check any local zip archive against the expected MD5."""
    archives = [path for path in files if path.suffix.lower() == ".zip"]
    print("\nARCHIVE MD5")
    if not archives:
        print(f"  No zip archive found under {root}; expected MD5 {EXPECTED_MD5} not checked.")
        return
    for archive in archives:
        digest = md5_file(archive)
        status = "MATCH" if digest.lower() == EXPECTED_MD5 else "MISMATCH"
        print(f"  {archive.relative_to(root)}: {digest} ({status})")


def inspect_workbooks(files: list[Path]) -> None:
    """Report workbook sheets, headers, and sample rows to identify demographic fields."""
    workbooks = [
        path for path in files if path.suffix.lower() in {".xlsx", ".xls", ".csv"}
    ]
    print("\nDEMOGRAPHICS / ASSESSMENT TABLES")
    if not workbooks:
        print("  No .xlsx, .xls, or .csv table found.")
        return

    for path in workbooks:
        print(f"  Table: {path}")
        if path.suffix.lower() == ".csv":
            try:
                import csv

                with path.open(encoding="utf-8-sig", newline="") as stream:
                    rows = list(csv.reader(stream))
                print(f"    Columns: {rows[0] if rows else []}")
                for row in rows[1:4]:
                    print(f"    Sample: {row}")
            except (OSError, UnicodeError, csv.Error) as exc:
                print(f"    Could not read CSV: {exc}")
            continue

        try:
            from openpyxl import load_workbook
        except ImportError:
            print("    openpyxl unavailable; reading workbook XML with the standard library.")
            try:
                for sheet_name, rows in read_xlsx_rows(path):
                    nonempty = [row for row in rows if any(value != "" for value in row)]
                    header = nonempty[0] if nonempty else []
                    print(f"    Sheet {sheet_name!r}: columns={header}")
                    for sample in nonempty[1:4]:
                        print(f"      Sample: {sample}")
                    group_column = next(
                        (
                            index
                            for index, name in enumerate(header)
                            if re.search(r"group|diagnos|class|^type$", name, re.IGNORECASE)
                        ),
                        None,
                    )
                    if group_column is not None:
                        groups = Counter(
                            row[group_column]
                            for row in nonempty[1:]
                            if group_column < len(row) and row[group_column]
                        )
                        print(f"      Group counts from {header[group_column]!r}: {dict(groups)}")
            except (OSError, KeyError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
                print(f"    Could not read workbook XML: {exc}")
            continue

        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                samples = [next(rows, ()) for _ in range(3)]
                print(f"    Sheet {sheet.title!r}: columns={list(header)}")
                for sample in samples:
                    if any(value is not None for value in sample):
                        print(f"      Sample: {list(sample)}")
            workbook.close()
        except (OSError, ValueError, KeyError) as exc:
            print(f"    Could not read workbook: {exc}")


def read_xlsx_rows(path: Path) -> list[tuple[str, list[list[str]]]]:
    """Read sheet cell text from an xlsx file using only the standard library."""
    namespace = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "package": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    with zipfile.ZipFile(path) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall("main:si", namespace):
                shared_strings.append(
                    "".join(text.text or "" for text in item.iter(f"{{{namespace['main']}}}t"))
                )

        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            relation.attrib["Id"]: relation.attrib["Target"]
            for relation in relationships.findall("package:Relationship", namespace)
        }
        sheets = []
        for sheet in workbook.findall("main:sheets/main:sheet", namespace):
            rel_id = sheet.attrib[f"{{{namespace['rel']}}}id"]
            target = targets[rel_id].lstrip("/")
            member = target if target.startswith("xl/") else f"xl/{target}"
            root = ET.fromstring(archive.read(member))
            rows: list[list[str]] = []
            for row in root.findall(".//main:sheetData/main:row", namespace):
                values: list[str] = []
                for cell in row.findall("main:c", namespace):
                    reference = cell.attrib.get("r", "")
                    column_letters = re.match(r"[A-Z]+", reference)
                    if column_letters is None:
                        continue
                    column = 0
                    for letter in column_letters.group():
                        column = column * 26 + ord(letter) - ord("A") + 1
                    while len(values) < column:
                        values.append("")
                    value = cell.findtext("main:v", default="", namespaces=namespace)
                    if cell.attrib.get("t") == "s" and value:
                        value = shared_strings[int(value)]
                    elif cell.attrib.get("t") == "inlineStr":
                        value = "".join(
                            text.text or ""
                            for text in cell.iter(f"{{{namespace['main']}}}t")
                        )
                    values[column - 1] = value
                rows.append(values)
            sheets.append((sheet.attrib.get("name", ""), rows))
    return sheets


def inspect_mat_file(path: Path) -> None:
    """Report MAT variables/shapes and available acquisition metadata."""
    print(f"\nMAT FILE DETAILS: {path}")
    try:
        import h5py

        if h5py.is_hdf5(path):
            print("  MAT storage: HDF5 (likely MATLAB v7.3)")
            with h5py.File(path, "r") as mat_file:
                def visit(name: str, item: Any) -> None:
                    if isinstance(item, h5py.Dataset):
                        print(f"  dataset {name}: shape={item.shape}, dtype={item.dtype}")
                        if item.size <= 32 and item.dtype.kind in "iuf":
                            try:
                                print(f"    values={item[()].tolist()}")
                            except (OSError, TypeError, ValueError):
                                pass
                    for key, value in getattr(item, "attrs", {}).items():
                        if key.lower() in {
                            "fs", "srate", "samplingrate", "sampling_rate",
                            "units", "unit", "reference", "ref",
                        }:
                            print(f"    attribute {name}/{key}: {value}")

                mat_file.visititems(visit)
            return
    except ImportError:
        pass

    try:
        from scipy.io import loadmat, whosmat

        descriptions = whosmat(path)
        for name, shape, matlab_type in descriptions:
            print(f"  variable {name}: shape={shape}, class={matlab_type}")
        signal_candidates = [
            (name, shape)
            for name, shape, _ in descriptions
            if len(shape) == 2 and min(shape, default=0) >= 32
        ]
        metadata_names = [
            name
            for name, _, _ in descriptions
            if re.search(r"sampling|srate|^fs$|channel|label|unit|refer", name, re.IGNORECASE)
        ]
        if signal_candidates:
            signal_name, shape = max(
                signal_candidates, key=lambda item: item[1][0] * item[1][1]
            )
            metadata_names.append(signal_name)
        values = loadmat(
            path,
            variable_names=list(dict.fromkeys(metadata_names)),
            squeeze_me=True,
            struct_as_record=False,
        )
        for name, value in values.items():
            if name.startswith("__"):
                continue
            fields = getattr(value, "_fieldnames", None)
            if fields:
                print(f"  struct {name} fields: {fields}")
            elif getattr(value, "size", 1) <= 32:
                print(f"  metadata {name}: {value}")

        for name, shape in signal_candidates:
            if name not in values:
                continue
            signal = values[name]
            numeric = getattr(signal, "dtype", None)
            print(
                f"  candidate signal {name}: shape={shape}, dtype={numeric}, "
                f"min={signal.min():.6g}, max={signal.max():.6g}"
            )
            fs_value = next(
                (
                    float(value.squeeze())
                    for key, value in values.items()
                    if re.search(r"sampling|srate|^fs$", key, re.IGNORECASE)
                    and getattr(value, "size", 0) == 1
                ),
                None,
            )
            if fs_value:
                samples = max(shape)
                print(
                    f"  estimated duration if [channels, samples] at {fs_value:g} Hz: "
                    f"{samples / fs_value:.3f} s"
                )
            break

        metadata_keys = {name.lower() for name, _, _ in descriptions}
        print(
            "  channel names/montage: "
            + ("metadata variable present; inspect it above." if any("chan" in key or "label" in key for key in metadata_keys)
               else "not encoded as a separate MAT variable.")
        )
        print(
            "  units/reference: "
            + ("metadata variable(s) present; inspect above." if any("unit" in key or "refer" in key for key in metadata_keys)
               else "not explicitly encoded as a separate MAT variable.")
        )
    except ImportError:
        print("  Neither h5py nor scipy is available to inspect MAT contents.")
    except (OSError, ValueError, NotImplementedError) as exc:
        print(f"  MAT contents could not be decoded: {exc}")


def main() -> int:
    """Print a local-only MODMA inspection summary."""
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="backslashreplace")

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--root",
        type=Path,
        help="Dataset directory; defaults to data/raw/modma or a detected local MODMA folder.",
    )
    args = parser.parse_args()

    root = find_dataset_root(DATA_ROOT, args.root)
    files = sorted(path for path in root.rglob("*") if path.is_file())
    print("MODMA INSPECTION SUMMARY")
    print(f"  DATA_ROOT: {DATA_ROOT}")
    print(f"  Dataset root: {root}")
    print(f"  Files found: {len(files)}")
    print_file_inventory(root, files)
    print_archive_checks(root, files)
    inspect_workbooks(files)

    mat_files = [path for path in files if path.suffix.lower() == ".mat"]
    if mat_files:
        inspect_mat_file(mat_files[0])
    else:
        print("\nMAT FILE DETAILS\n  No .mat recording files found.")

    print("\nReview this report before implementing the loader or preprocessing pipeline.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
