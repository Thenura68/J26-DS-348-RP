"""Loading for the 128-channel resting-state MODMA EEG dataset."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from scipy.io import loadmat, whosmat

from snn_depression.common.paths import DATA_ROOT

_DATASET_DIRECTORY = "EEG_128channels_resting_lanzhou_2015"
_SUBJECT_ID_PATTERN = re.compile(r"^(\d+)")
_EXPECTED_SUBJECTS = 53
_EXPECTED_GROUP_COUNTS = {"MDD": 24, "HC": 29}


@dataclass
class SubjectRecording:
    """One subject's recording; EEG is [channels, samples] in native MAT units."""

    subject_id: str
    eeg: np.ndarray
    fs: float
    channel_names: list[str]
    phq9: int | None
    group: str
    meta: dict[str, Any] = field(default_factory=dict)


def _find_dataset_root() -> Path:
    """Find the dataset under DATA_ROOT without relying on a machine-specific path."""
    candidates = (
        DATA_ROOT / "raw" / "modma",
        DATA_ROOT / _DATASET_DIRECTORY,
        DATA_ROOT,
    )
    for candidate in candidates:
        if candidate.is_dir() and any(
            "subjects_information" in path.name.lower()
            for path in candidate.rglob("*.xlsx")
        ):
            return candidate
    tried = ", ".join(str(path) for path in candidates)
    raise FileNotFoundError(
        "Could not find the MODMA subject-information workbook. "
        f"Checked dataset roots under DATA_ROOT={DATA_ROOT}: {tried}"
    )


def _subject_id(value: object, source: str) -> str:
    """Normalize a workbook/file identifier to the eight-digit MODMA subject ID."""
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Invalid MODMA subject ID {value!r} in {source}")
    if isinstance(value, (int, np.integer)):
        normalized = f"{int(value):08d}"
    elif isinstance(value, (float, np.floating)) and float(value).is_integer():
        normalized = f"{int(value):08d}"
    else:
        text = str(value).strip()
        match = _SUBJECT_ID_PATTERN.match(text)
        if match is None:
            raise ValueError(f"Invalid MODMA subject ID {value!r} in {source}")
        normalized = match.group(1)
    return normalized


def _read_subject_table(path: Path) -> dict[str, dict[str, Any]]:
    """Read demographics and PHQ-9 values keyed by the workbook subject ID."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise ValueError(f"MODMA subject-information workbook is empty: {path}")
        headers = [
            str(value).strip() if value is not None else ""
            for value in header_row
        ]
        normalized_headers = [header.casefold() for header in headers]
        required = {"subject id", "type", "phq-9"}
        missing = required.difference(normalized_headers)
        if missing:
            raise ValueError(
                f"MODMA workbook {path} is missing required column(s): "
                f"{', '.join(sorted(missing))}; found {headers}"
            )

        subject_column = normalized_headers.index("subject id")
        group_column = normalized_headers.index("type")
        phq9_column = normalized_headers.index("phq-9")
        subjects: dict[str, dict[str, Any]] = {}
        for row_number, values in enumerate(rows, start=2):
            if subject_column >= len(values) or values[subject_column] is None:
                continue
            subject_id = _subject_id(
                values[subject_column], f"{path} (row {row_number})"
            )
            if subject_id in subjects:
                raise ValueError(
                    f"Duplicate subject ID {subject_id} in MODMA workbook "
                    f"{path} (row {row_number})"
                )
            row = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            group = str(values[group_column]).strip().upper()
            if group not in _EXPECTED_GROUP_COUNTS:
                raise ValueError(
                    f"Invalid group {values[group_column]!r} for subject "
                    f"{subject_id} in {path} (row {row_number}); expected MDD or HC"
                )
            score = values[phq9_column] if phq9_column < len(values) else None
            if score is not None:
                try:
                    score_number = float(score)
                except (TypeError, ValueError) as exc:
                    raise ValueError(
                        f"Invalid PHQ-9 value {score!r} for subject {subject_id} "
                        f"in {path} (row {row_number})"
                    ) from exc
                if not score_number.is_integer():
                    raise ValueError(
                        f"Non-integer PHQ-9 value {score!r} for subject "
                        f"{subject_id} in {path} (row {row_number})"
                    )
                row["PHQ-9"] = int(score_number)
            row["type"] = group
            subjects[subject_id] = row
        return subjects
    finally:
        workbook.close()


def _recording_files(root: Path) -> dict[str, Path]:
    """Index the dataset's MAT recordings by their leading numeric subject IDs."""
    files: dict[str, Path] = {}
    for path in sorted(root.rglob("*.mat")):
        match = _SUBJECT_ID_PATTERN.match(path.stem)
        if match is None:
            continue
        subject_id = match.group(1)
        if subject_id in files:
            raise ValueError(
                f"Multiple resting-state MAT files found for subject {subject_id}: "
                f"{files[subject_id]} and {path}"
            )
        files[subject_id] = path
    return files


def _subject_index() -> tuple[Path, dict[str, Path], dict[str, dict[str, Any]]]:
    """Validate the file/workbook subject join and cohort counts."""
    root = _find_dataset_root()
    workbooks = [
        path
        for path in root.rglob("*.xlsx")
        if "subjects_information" in path.name.casefold()
    ]
    if len(workbooks) != 1:
        raise FileNotFoundError(
            f"Expected exactly one MODMA subjects_information workbook under {root}; "
            f"found {len(workbooks)}: {workbooks}"
        )

    recordings = _recording_files(root)
    subjects = _read_subject_table(workbooks[0])
    file_ids = set(recordings)
    table_ids = set(subjects)
    missing_rows = sorted(file_ids - table_ids)
    missing_recordings = sorted(table_ids - file_ids)
    if missing_rows or missing_recordings:
        raise ValueError(
            "MODMA subject join mismatch: "
            f"{len(recordings)} resting MAT files and {len(subjects)} workbook rows; "
            f"MAT IDs missing workbook rows={missing_rows}; "
            f"workbook IDs missing MAT files={missing_recordings}"
        )

    group_counts = {
        group: sum(row["type"] == group for row in subjects.values())
        for group in _EXPECTED_GROUP_COUNTS
    }
    if len(file_ids) != _EXPECTED_SUBJECTS or group_counts != _EXPECTED_GROUP_COUNTS:
        raise ValueError(
            "Unexpected MODMA cohort after subject-ID join: "
            f"matched={len(file_ids)}, groups={group_counts}; expected "
            f"matched={_EXPECTED_SUBJECTS}, groups={_EXPECTED_GROUP_COUNTS}"
        )
    return root, recordings, subjects


def list_subjects(montage: str = "eeg_128ch_resting") -> list[str]:
    """Return sorted subject IDs after validating every recording/workbook match."""
    if montage != "eeg_128ch_resting":
        raise ValueError(
            f"Unsupported MODMA montage {montage!r}; only 'eeg_128ch_resting' "
            "is available in this loader"
        )
    _, recordings, _ = _subject_index()
    return sorted(recordings)


def load_subject(
    subject_id: str, montage: str = "eeg_128ch_resting"
) -> SubjectRecording:
    """Load one resting-state recording with all 129 rows preserved.

    The local MODMA paper identifies rows E1--E128 followed by Cz, the reference
    electrode. It does not state the stored amplitude unit, so the returned values
    remain in native MAT units until that unit is confirmed.
    """
    if montage != "eeg_128ch_resting":
        raise ValueError(
            f"Unsupported MODMA montage {montage!r}; only 'eeg_128ch_resting' "
            "is available in this loader"
        )
    requested_id = _subject_id(subject_id, "load_subject argument")
    root, recordings, subjects = _subject_index()
    if requested_id not in recordings:
        raise KeyError(
            f"MODMA subject {requested_id} is not available under {root}; "
            f"available subject IDs: {', '.join(sorted(recordings))}"
        )

    workbook_row = subjects[requested_id]
    phq_value = workbook_row.get("PHQ-9")
    if phq_value is None:
        raise ValueError(
            f"MODMA subject {requested_id} has no PHQ-9 score in the "
            "subjects_information workbook"
        )

    path = recordings[requested_id]
    variables = whosmat(path)
    signal_candidates = [
        (name, shape)
        for name, shape, matlab_type in variables
        if matlab_type in {"double", "single"}
        and len(shape) == 2
        and shape[0] == 129
        and shape[1] > 129
    ]
    if len(signal_candidates) != 1:
        raise ValueError(
            f"Expected one [129, samples] EEG matrix in {path}; found "
            f"{[(name, shape) for name, shape in signal_candidates]}"
        )

    signal_name = signal_candidates[0][0]
    mat_data = loadmat(path, variable_names=[signal_name, "samplingRate"])
    eeg = np.asarray(mat_data[signal_name], dtype=np.float64)
    fs_value = np.asarray(mat_data.get("samplingRate", [])).squeeze()
    if fs_value.size != 1:
        raise ValueError(f"Missing or invalid samplingRate value in {path}")
    fs = float(fs_value)
    if not np.isfinite(fs) or fs <= 0:
        raise ValueError(f"Invalid sampling rate {fs!r} Hz in {path}")
    if eeg.ndim != 2 or eeg.shape[0] != 129 or not np.isfinite(eeg).all():
        raise ValueError(
            f"Invalid EEG matrix in {path}: expected finite [129, samples], "
            f"got shape={eeg.shape}"
        )

    channel_names = [f"E{index}" for index in range(1, 129)] + ["Cz"]
    metadata = {
        "age": workbook_row.get("age"),
        "gender": workbook_row.get("gender"),
        "demographics": workbook_row,
        "source_file": str(path.relative_to(root)),
        "signal_variable": signal_name,
        "unit": "unknown (native MAT values; not converted)",
        "unit_conversion_applied": False,
        "reference": "Cz recorded in the 129th row; retained",
        "reference_removed": False,
        "channel_name_basis": "MODMA paper: E1-E128, followed by Cz reference",
    }
    return SubjectRecording(
        subject_id=requested_id,
        eeg=eeg,
        fs=fs,
        channel_names=channel_names,
        phq9=int(phq_value),
        group=str(workbook_row["type"]),
        meta=metadata,
    )
